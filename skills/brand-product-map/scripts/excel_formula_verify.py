"""Open and sanity-check a generated workbook.

Exit codes:
  0  workbook opened and formulas are present without obvious formula errors
  2  input or workbook validation error
  4  workbook opened successfully but contains no formulas (normal for this skill)

This intentionally does not recalculate formulas. openpyxl can inspect formula
text, while recalculation requires a separate Excel/LibreOffice runtime.
"""
import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", "--workbook", dest="workbook", required=True)
    ap.add_argument("--rows", help="Optional normalized rows.json for row-count validation")
    ap.add_argument("--sheet", default="在售产品")
    args = ap.parse_args()

    path = Path(args.workbook)
    if not path.is_file():
        print(f"ERROR: workbook not found: {path}", file=sys.stderr)
        return 2
    try:
        wb = load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:
        print(f"ERROR: workbook cannot be opened: {exc}", file=sys.stderr)
        return 2

    if args.sheet not in wb.sheetnames:
        print(f"ERROR: sheet not found: {args.sheet}; sheets={wb.sheetnames}", file=sys.stderr)
        return 2
    ws = wb[args.sheet]
    data_rows = max(ws.max_row - 1, 0)
    if args.rows:
        try:
            expected = len(json.loads(Path(args.rows).read_text(encoding="utf-8-sig")))
        except Exception as exc:
            print(f"ERROR: rows file cannot be read: {exc}", file=sys.stderr)
            return 2
        if data_rows != expected:
            print(f"ERROR: row count mismatch: workbook={data_rows} rows.json={expected}", file=sys.stderr)
            return 2

    formulas = []
    formula_errors = []
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                formulas.append(cell.coordinate)
                if any(token in value.upper() for token in ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!")):
                    formula_errors.append(f"{cell.coordinate}:{value}")
    if formula_errors:
        print("ERROR: obvious formula errors: " + "; ".join(formula_errors), file=sys.stderr)
        return 2
    if not formulas:
        print(f"OK: workbook opened; rows={data_rows}; formulas=0 (normal exit 4)")
        return 4
    print(f"OK: workbook opened; rows={data_rows}; formulas={len(formulas)}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        raise SystemExit(130)
