"""Build the standardized brand-product xlsx from normalized rows.

Usage:
  python build_excel.py --rows rows.json --out "<Brand>在售产品整理.xlsx" --brand "Brand" --excluded '[{"status":"待核实","count":1,"note":"..."}]'

rows.json schema (list):
  {
    "name": "Eye Mask", "capacity": "60 Pairs/Pcs", "pack": 2,
    "price": 16.99, "monthly_sales": 3000, "monthly_revenue": 50970.0,
    "rating": 4.3, "reviews": 1234, "category": "眼膜", "asin": "B0XXX",
    "url": "https://..."
  }
"""
import argparse, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER = ["序号","产品名称","容量规格","包装","价格$","SS月销量","SS月销售额$","评分","评论数","品类","ASIN","商品链接"]

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--rows", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--brand", required=True)
    ap.add_argument("--excluded", default="[]", help="JSON list of excluded statuses")
    args = ap.parse_args()

    rows = json.load(open(args.rows, encoding="utf-8-sig"))
    excluded = json.loads(args.excluded)

    wb = Workbook()
    ws = wb.active
    ws.title = "在售产品"

    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c, h in enumerate(HEADER, 1):
        cell = ws.cell(1, c, h)
        cell.fill = head_fill; cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for i, r in enumerate(rows, 1):
        pack = f'{r["pack"]} Pack' if r.get("pack") and r["pack"] > 1 else ""
        vals = [
            i,
            r.get("name",""),
            r.get("capacity",""),
            pack,
            r.get("price"),
            r.get("monthly_sales"),
            r.get("monthly_revenue"),
            r.get("rating"),
            r.get("reviews"),
            r.get("category",""),
            r.get("asin",""),
            r.get("url",""),
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(i+1, c, v if v != "" else None)
            cell.border = border
            if c in (5,7) and isinstance(v,(int,float)):
                cell.number_format = '"$"#,##0.00'
            if c == 12 and v:
                cell.hyperlink = v; cell.font = Font(color="0563C1", underline="single")

    widths = [6,32,18,10,10,12,14,8,10,16,14,40]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"

    # 口径说明 sheet
    ws2 = wb.create_sheet("口径说明")
    notes = [
        [f"品牌：{args.brand}"],
        [f"在售条数：{len(rows)}"],
        ["数据来源：飞书 Base 子体表，状态=在售"],
        ["同产品不同容量/包装已拆为多行；缺失销量/销售额留空，不填 0"],
        ["包装数从变体属性正则提取 Pack of N / N Pack；含 '+' 视为 2 Pack 组合套装"],
        [""],
        ["被剔除状态："],
    ]
    for ex in excluded:
        notes.append([f"  - {ex.get('status','?')}：{ex.get('count',0)} 条 — {ex.get('note','')}"])
    for r_i, row in enumerate(notes, 1):
        ws2.cell(r_i, 1, row[0])
    ws2.column_dimensions["A"].width = 80

    wb.save(args.out)
    print(f"wrote {args.out} rows={len(rows)}")

if __name__ == "__main__":
    main()
